"""Étape 2 - Contrôle qualité et tableau de bord.

Interroge les vues du schéma controle et la vue des indicateurs communaux,
puis produit un classeur Excel : synthèse des anomalies, détail par contrôle
(avec coordonnées pour les retrouver), indicateurs par commune.
"""
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config
from outils import log, pg_lire

CONTROLES = {
    "Géométries invalides": "controle.v_geometrie_invalide",
    "Chevauchements de zones": "controle.v_zonage_chevauchement",
    "Parcelles hors zonage": "controle.v_parcelle_hors_zonage",
    "Adresses en doublon": "controle.v_adresse_doublon",
    "Adresses non certifiées": "controle.v_adresse_non_certifiee",
    "Logements sans adresse": "controle.v_batiment_sans_adresse",
    "Arbres en doublon": "controle.v_arbre_doublon",
    "Fiches arbre incomplètes": "controle.v_arbre_incomplet",
}

ACTIONS = {
    "Géométries invalides": "Corriger la géométrie (Vérifier la validité / Réparer les géométries dans QGIS)",
    "Chevauchements de zones": "Signaler à la commune : à corriger lors de la prochaine procédure du PLU",
    "Parcelles hors zonage": "Vérifier le plan de zonage déposé au GPU (trou de numérisation ou domaine public)",
    "Adresses en doublon": "Supprimer ou renuméroter dans la Base Adresse Locale",
    "Adresses non certifiées": "Accompagner la commune pour certifier ses adresses (mes-adresses.data.gouv.fr)",
    "Logements sans adresse": "Vérifier sur le terrain puis créer l'adresse dans la Base Adresse Locale",
    "Arbres en doublon": "Fusionner les deux fiches (garder l'historique le plus complet)",
    "Fiches arbre incomplètes": "Compléter lors de la prochaine tournée QField",
}

ENTETE = PatternFill("solid", fgColor="1F4E5A")


def mettre_en_forme(ws, largeurs=None):
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = ENTETE
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for i, col in enumerate(ws.columns, start=1):
        largeur = max(len(str(c.value or "")) for c in list(col)[:200])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, largeur + 2), 70)
    for lettre, l in (largeurs or {}).items():
        ws.column_dimensions[lettre].width = l


def main():
    sortie = config.LIVRABLES / "08_Tableaux" / "controle_qualite_et_indicateurs.xlsx"
    sortie.parent.mkdir(parents=True, exist_ok=True)

    synthese, details = [], {}
    for nom, vue in CONTROLES.items():
        gdf = pg_lire(f"SELECT * FROM {vue}")
        if len(gdf):
            gdf["x_l93"] = gdf.geometry.x.round(1)
            gdf["y_l93"] = gdf.geometry.y.round(1)
        details[nom] = pd.DataFrame(gdf.drop(columns="geom"))
        synthese.append({"Contrôle": nom, "Vue PostGIS": vue, "Anomalies": len(gdf),
                         "Action corrective": ACTIONS[nom]})
        log(f"  {nom} : {len(gdf)}")

    ind = pg_lire("SELECT * FROM ref_admin.mv_indicateurs_commune", geom="geom")
    ind = pd.DataFrame(ind.drop(columns="geom"))

    with pd.ExcelWriter(sortie, engine="openpyxl") as xl:
        pd.DataFrame(synthese).to_excel(xl, sheet_name="Synthèse contrôles", index=False)
        ind.to_excel(xl, sheet_name="Indicateurs communes", index=False)
        for nom, df in details.items():
            (df if len(df) else pd.DataFrame({"résultat": ["Aucune anomalie"]})).to_excel(
                xl, sheet_name=nom[:31], index=False)
        for ws in xl.book.worksheets:
            mettre_en_forme(ws, {"D": 80} if ws.title == "Synthèse contrôles" else None)
    log(f"  -> {sortie.relative_to(config.ROOT)}")


if __name__ == "__main__":
    main()
